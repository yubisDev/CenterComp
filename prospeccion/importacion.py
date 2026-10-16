"""Importación unificada de leads desde CSV o Excel.

Acepta archivos armados a mano y el CSV que descarga Meta Lead Ads (que
viene en UTF-16, separado por tabulaciones y con columnas en inglés como
full_name, phone_number o campaign_name). Si un lead ya existe —mismo
teléfono, mismo correo o, sin ninguno de los dos, mismo nombre— no se
duplica: se completa el registro existente con lo nuevo."""
import csv
import io
import re
import unicodedata

from django.db import transaction
from django.utils import timezone

from .models import CategoriaInteres, Comprador
from .normalizacion import normalizar_pais, normalizar_telefono
from .texto import titulo_inteligente

TAMANO_MAXIMO = 10 * 1024 * 1024  # 10 MB


def normalizar_clave(texto):
    """'¿Qué te interesa?' -> 'que_te_interesa' (para comparar encabezados y valores)."""
    texto = unicodedata.normalize('NFKD', str(texto or '').strip().lower())
    texto = ''.join(ch for ch in texto if not unicodedata.combining(ch))
    texto = re.sub(r'[^a-z0-9]+', '_', texto)
    return texto.strip('_')


# Campo del modelo -> encabezados que se aceptan (ya normalizados).
# El orden importa: si el archivo trae varias, gana la primera que tenga valor.
COLUMNAS = {
    'nombre_empresa': ['nombre_empresa', 'empresa', 'company', 'company_name', 'razon_social'],
    'nombre_contacto': ['nombre_contacto', 'nombre', 'full_name', 'nombre_completo', 'name', 'contacto'],
    'first_name': ['first_name', 'nombres'],
    'last_name': ['last_name', 'apellidos', 'apellido'],
    'email': ['email', 'correo', 'correo_electronico', 'e_mail', 'mail'],
    'telefono': ['telefono', 'phone_number', 'phone', 'celular', 'movil', 'numero', 'telefono_whatsapp'],
    'whatsapp': ['whatsapp', 'whatsapp_number', 'numero_whatsapp'],
    'pais': ['pais', 'country'],
    'ciudad': ['ciudad', 'city'],
    'sector': ['sector', 'industria', 'industry'],
    'linkedin_url': ['linkedin', 'linkedin_url'],
    'facebook_url': ['facebook', 'facebook_url'],
    'instagram_url': ['instagram', 'instagram_url'],
    'sitio_web': ['sitio_web', 'web', 'website', 'pagina_web'],
    'fuente': ['fuente', 'origen', 'canal', 'source', 'platform'],
    'campana': ['campana', 'campaign_name', 'campaign', 'ad_name', 'anuncio', 'form_name', 'adset_name'],
    'categoria': ['categoria', 'categorias', 'categoria_interes', 'interes', 'intereses', 'le_interesa'],
    'palabras_clave_interes': ['palabras_clave', 'palabras_clave_interes'],
    'notas': ['notas', 'nota', 'comentarios', 'observaciones', 'mensaje'],
    'acepta_mensajes': ['acepta_mensajes', 'acepta', 'opt_in', 'autoriza', 'acepta_promociones'],
}

# Columnas que delatan un export de Meta Lead Ads.
COLUMNAS_META = {'ad_id', 'ad_name', 'adset_id', 'campaign_id', 'form_id', 'is_organic', 'platform'}

FUENTE_ALIASES = {
    'hunter': Comprador.Fuente.HUNTER,
    'hunter_io': Comprador.Fuente.HUNTER,
    'apollo': Comprador.Fuente.APOLLO,
    'procolombia': Comprador.Fuente.PROCOLOMBIA,
    'camara_de_comercio': Comprador.Fuente.CAMARA_COMERCIO,
    'camara_comercio': Comprador.Fuente.CAMARA_COMERCIO,
    'linkedin': Comprador.Fuente.LINKEDIN,
    'linkedin_sales_navigator': Comprador.Fuente.LINKEDIN,
    'manual': Comprador.Fuente.MANUAL,
    'whatsapp': Comprador.Fuente.WHATSAPP,
    'whatsapp_business': Comprador.Fuente.WHATSAPP,
    'wa': Comprador.Fuente.WHATSAPP,
    'instagram': Comprador.Fuente.INSTAGRAM,
    'instagram_dm': Comprador.Fuente.INSTAGRAM,
    'ig': Comprador.Fuente.INSTAGRAM_ADS,  # "platform" en Meta Lead Ads
    'instagram_ads': Comprador.Fuente.INSTAGRAM_ADS,
    'fb': Comprador.Fuente.FACEBOOK_ADS,  # "platform" en Meta Lead Ads
    'facebook': Comprador.Fuente.FACEBOOK_ADS,
    'facebook_ads': Comprador.Fuente.FACEBOOK_ADS,
    'meta': Comprador.Fuente.META_ADS,
    'meta_ads': Comprador.Fuente.META_ADS,
    'anuncio': Comprador.Fuente.META_ADS,
    'formulario': Comprador.Fuente.META_ADS,
    'web': Comprador.Fuente.PAGINA_WEB,
    'pagina_web': Comprador.Fuente.PAGINA_WEB,
    'sitio_web': Comprador.Fuente.PAGINA_WEB,
    'pagina': Comprador.Fuente.PAGINA_WEB,
    'ia': Comprador.Fuente.IA,
}
for _valor, _ in Comprador.Fuente.choices:
    FUENTE_ALIASES.setdefault(_valor, _valor)

VALORES_SI = {'si', 'yes', 'true', '1', 'x', 'acepta', 'acepto'}

CAMPOS_A_COMPLETAR = [
    'nombre_empresa', 'nombre_contacto', 'email', 'telefono', 'whatsapp', 'pais', 'ciudad', 'sector',
    'linkedin_url', 'facebook_url', 'instagram_url', 'sitio_web', 'campana',
]


# ------------------------------------------------------------------ lectura

def _a_texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # Excel guarda los teléfonos como 3001234567.0
    return str(valor).strip()


def _leer_csv(crudo):
    if crudo.startswith((b'\xff\xfe', b'\xfe\xff')):
        texto = crudo.decode('utf-16')  # así exporta Meta Lead Ads
    else:
        try:
            texto = crudo.decode('utf-8-sig')
        except UnicodeDecodeError:
            texto = crudo.decode('latin-1')  # Excel en Windows guardando "CSV"
    muestra = texto[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=',;\t')
    except csv.Error:
        dialecto = csv.excel
    lector = csv.reader(io.StringIO(texto), dialecto)
    filas = list(lector)
    if not filas:
        return [], []
    return filas[0], filas[1:]


def _leer_excel(archivo):
    import openpyxl

    wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    filas = list(wb.active.iter_rows(values_only=True))
    if not filas:
        return [], []
    return filas[0], filas[1:]


def leer_archivo(archivo):
    """Devuelve (filas, error). Cada fila es un dict con encabezados normalizados."""
    if archivo.size > TAMANO_MAXIMO:
        return [], 'El archivo supera el tamaño máximo permitido (10 MB).'
    nombre = archivo.name.lower()
    if nombre.endswith(('.csv', '.txt', '.tsv')):
        encabezados, datos = _leer_csv(archivo.read())
    elif nombre.endswith(('.xlsx', '.xlsm')):
        encabezados, datos = _leer_excel(archivo)
    else:
        return [], 'Formato de archivo no soportado. Usa CSV o Excel (.xlsx).'

    claves = [normalizar_clave(h) for h in encabezados]
    filas = []
    for fila in datos:
        valores = {}
        for clave, valor in zip(claves, fila):
            if clave and clave not in valores:
                valores[clave] = _a_texto(valor)
        if any(valores.values()):
            filas.append(valores)
    return filas, None


def _columnas_de_interes(claves):
    """Preguntas personalizadas de un formulario de Meta como
    '¿Qué producto te interesa?' también cuentan como categoría."""
    return [c for c in claves if 'interes' in c and c not in COLUMNAS['palabras_clave_interes']]


def mapear_fila(fila):
    """Pasa una fila con encabezados libres a los campos del modelo."""
    datos = {}
    for campo, alias in COLUMNAS.items():
        for clave in alias:
            if fila.get(clave):
                datos[campo] = fila[clave]
                break
    if not datos.get('nombre_contacto') and (datos.get('first_name') or datos.get('last_name')):
        datos['nombre_contacto'] = f"{datos.get('first_name', '')} {datos.get('last_name', '')}".strip()
    datos.pop('first_name', None)
    datos.pop('last_name', None)

    textos_interes = [fila[c] for c in _columnas_de_interes(fila.keys()) if fila.get(c)]
    if datos.get('categoria'):
        textos_interes.insert(0, datos['categoria'])
    datos['categoria'] = ' | '.join(dict.fromkeys(textos_interes))
    return datos


# -------------------------------------------------------------- categorías

def detectar_categorias(textos, categorias):
    """Categorías cuyo nombre o palabras clave aparecen en alguno de los textos."""
    texto = ' '.join(normalizar_clave(t) for t in textos if t)
    if not texto:
        return []
    texto = f'_{texto}_'
    encontradas = []
    for categoria in categorias:
        candidatos = [categoria.nombre] + categoria.lista_palabras()
        for palabra in candidatos:
            clave = normalizar_clave(palabra)
            # Coincidencia por inicio de palabra: "vehiculo" encuentra "vehiculos".
            if clave and f'_{clave}' in texto:
                encontradas.append(categoria)
                break
    return encontradas


# -------------------------------------------------------------- duplicados

def buscar_existente(datos, telefono_normalizado):
    if telefono_normalizado and len(telefono_normalizado) >= 8:
        existente = Comprador.objects.filter(telefono_normalizado=telefono_normalizado).first()
        if existente:
            return existente
    email = datos.get('email', '')
    if email:
        existente = Comprador.objects.filter(email__iexact=email).first()
        if existente:
            return existente
    if not telefono_normalizado and not email:
        filtro = {}
        if datos.get('nombre_empresa'):
            filtro['nombre_empresa__iexact'] = datos['nombre_empresa']
        if datos.get('nombre_contacto'):
            filtro['nombre_contacto__iexact'] = datos['nombre_contacto']
        if filtro:
            return Comprador.objects.filter(**filtro).first()
    return None


# ------------------------------------------------------------------ proceso

def procesar_importacion(archivo, fuente_defecto='', campana_defecto='', categorias_defecto=(),
                         acepta_defecto=False):
    """Importa el archivo y devuelve un resumen:
    {'creados': int, 'actualizados': int, 'sin_cambios': int, 'errores': [str], 'meta': bool}"""
    resumen = {'creados': 0, 'actualizados': 0, 'sin_cambios': 0, 'errores': [], 'meta': False}
    filas, error = leer_archivo(archivo)
    if error:
        resumen['errores'].append(error)
        return resumen
    if not filas:
        resumen['errores'].append('El archivo no tiene filas con datos.')
        return resumen

    es_meta = bool(COLUMNAS_META & set(filas[0].keys()))
    resumen['meta'] = es_meta
    categorias = list(CategoriaInteres.objects.filter(activa=True))
    categorias_defecto = list(categorias_defecto)

    for i, fila in enumerate(filas, start=2):
        try:
            with transaction.atomic():
                resultado = _importar_fila(
                    mapear_fila(fila), categorias, es_meta,
                    fuente_defecto, campana_defecto, categorias_defecto, acepta_defecto,
                )
        except Exception as exc:  # una fila rota no debe tumbar toda la importación
            resumen['errores'].append(f'Fila {i}: no se pudo guardar ({exc}).')
            continue
        if isinstance(resultado, str):
            resumen['errores'].append(f'Fila {i}: {resultado}')
        else:
            resumen[resultado[0]] += 1
    return resumen


def _importar_fila(datos, categorias, es_meta, fuente_defecto, campana_defecto, categorias_defecto,
                   acepta_defecto):
    for campo in ('nombre_empresa', 'nombre_contacto'):
        if datos.get(campo):
            datos[campo] = titulo_inteligente(re.sub(r'\s+', ' ', datos[campo]))[:200 if campo == 'nombre_empresa' else 150]
    numero = datos.get('telefono') or datos.get('whatsapp') or ''
    if not (datos.get('nombre_empresa') or datos.get('nombre_contacto') or numero):
        return 'no tiene nombre ni teléfono, se omitió.'

    fuente = FUENTE_ALIASES.get(normalizar_clave(datos.get('fuente', '')))
    if not fuente:
        fuente = fuente_defecto or (Comprador.Fuente.META_ADS if es_meta else Comprador.Fuente.MANUAL)
    campana = (datos.get('campana') or campana_defecto)[:200]
    pais = normalizar_pais(datos.get('pais', ''))
    acepta = normalizar_clave(datos.get('acepta_mensajes', '')) in VALORES_SI or acepta_defecto

    cats = detectar_categorias([datos.get('categoria', ''), campana], categorias)
    for c in categorias_defecto:
        if c not in cats:
            cats.append(c)

    valores = {
        'nombre_empresa': datos.get('nombre_empresa', ''),
        'nombre_contacto': datos.get('nombre_contacto', ''),
        'email': datos.get('email', '')[:254],
        'telefono': datos.get('telefono', '')[:30],
        'whatsapp': datos.get('whatsapp', '')[:30],
        'pais': pais,
        'ciudad': datos.get('ciudad', '')[:100],
        'sector': datos.get('sector', '')[:120],
        'linkedin_url': datos.get('linkedin_url', ''),
        'facebook_url': datos.get('facebook_url', ''),
        'instagram_url': datos.get('instagram_url', ''),
        'sitio_web': datos.get('sitio_web', ''),
        'campana': campana,
    }
    texto_interes = datos.get('categoria', '')
    notas = datos.get('notas', '')
    if texto_interes and not cats:
        # No coincidió con ninguna categoría: que no se pierda lo que escribió.
        notas = f'{notas}\nInterés indicado: {texto_interes}'.strip()

    existente = buscar_existente(valores, normalizar_telefono(numero, pais) if numero else '')
    if existente:
        return _fusionar(existente, valores, cats, acepta, fuente, notas)

    comprador = Comprador.objects.create(
        fuente=fuente, acepta_mensajes=acepta, notas=notas,
        palabras_clave_interes=datos.get('palabras_clave_interes', '')[:300],
        **valores,
    )
    if cats:
        comprador.categorias_interes.add(*cats)
    return ('creados',)


def _fusionar(existente, valores, cats, acepta, fuente, notas):
    cambio = False
    for campo in CAMPOS_A_COMPLETAR:
        if valores.get(campo) and not getattr(existente, campo):
            setattr(existente, campo, valores[campo])
            cambio = True
    if acepta and not existente.acepta_mensajes:
        existente.acepta_mensajes = True
        cambio = True

    if fuente != existente.fuente or (valores['campana'] and valores['campana'] != existente.campana):
        nuevo_origen = dict(Comprador.Fuente.choices).get(fuente, fuente)
        detalle = f' ({valores["campana"]})' if valores['campana'] else ''
        linea = f'{timezone.localdate():%d/%m/%Y}: volvió a llegar por {nuevo_origen}{detalle}.'
        if linea not in existente.notas:
            existente.notas = f'{existente.notas}\n{linea}'.strip()
            cambio = True
    if notas and notas not in existente.notas:
        existente.notas = f'{existente.notas}\n{notas}'.strip()
        cambio = True

    nuevas = [c for c in cats if not existente.categorias_interes.filter(pk=c.pk).exists()]
    if nuevas:
        existente.categorias_interes.add(*nuevas)
        cambio = True
    if cambio:
        existente.save()
        return ('actualizados',)
    return ('sin_cambios',)


# ---------------------------------------------------------- plantilla Excel

ENCABEZADOS_PLANTILLA = [
    'nombre_contacto', 'nombre_empresa', 'telefono', 'email', 'pais', 'ciudad',
    'fuente', 'campana', 'categoria', 'acepta_mensajes', 'notas',
]
EJEMPLOS_PLANTILLA = [
    ['Ana Pérez', '', '+507 6123 4567', 'ana@ejemplo.com', 'Panamá', 'Ciudad de Panamá',
     'whatsapp', 'Revestimientos Panamá', 'revestimientos', 'si', 'Pidió precios del porcelanato'],
    ['Carlos Ruiz', 'Transportes CR SAS', '300 123 4567', '', 'Colombia', 'Bogotá',
     'instagram', '', 'vehículos, inmuebles', 'no', ''],
]


def construir_plantilla_excel():
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Leads'
    ws.append(ENCABEZADOS_PLANTILLA)
    for fila in EJEMPLOS_PLANTILLA:
        ws.append(fila)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for columna in ws.columns:
        ws.column_dimensions[columna[0].column_letter].width = 22

    ayuda = wb.create_sheet('Instrucciones')
    lineas = [
        'Borra las filas de ejemplo antes de importar.',
        'Cada fila necesita al menos un nombre (persona o empresa) o un teléfono.',
        'telefono: con o sin indicativo. Sin indicativo se asume Colombia (10 dígitos) o Panamá (8 dígitos).',
        'fuente: whatsapp, instagram, meta_ads, facebook_ads, instagram_ads, pagina_web, manual...',
        'categoria: texto libre; se reconoce por palabras clave. Categorías: '
        + ', '.join(CategoriaInteres.objects.filter(activa=True).values_list('nombre', flat=True)),
        'acepta_mensajes: si / no. Solo "si" cuando la persona aceptó recibir promociones.',
        'Si el teléfono o el correo ya existen, no se duplica: se completa el registro existente.',
        'También puedes subir directamente el CSV que descarga Meta Lead Ads.',
    ]
    for linea in lineas:
        ayuda.append([linea])
    ayuda.column_dimensions['A'].width = 110

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
