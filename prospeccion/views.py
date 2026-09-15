import csv
import io
import threading
from decimal import Decimal, InvalidOperation
from urllib.parse import quote

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.mail import EmailMessage, get_connection
from django.core.paginator import Paginator
from django.db.models import Count, Q, Sum
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from . import ia_busqueda
from .forms import (
    BusquedaIAForm,
    CompradorForm,
    HistorialContactoForm,
    HistorialContactoProveedorForm,
    ImportarCompradoresForm,
    ImportarProductosForm,
    PlantillaMensajeForm,
    ProductoForm,
    ProveedorForm,
)
from .models import (
    BusquedaIA,
    Comprador,
    HistorialContacto,
    HistorialContactoProveedor,
    PlantillaMensaje,
    Producto,
    Proveedor,
)
from .texto import titulo_inteligente

_SIN_CARGAR = object()  # centinela: "no me pasaron la plantilla, búscala tú" (distinto de None = "ya sé que no hay")

FUENTE_ALIASES = {
    'hunter': Comprador.Fuente.HUNTER,
    'hunter.io': Comprador.Fuente.HUNTER,
    'apollo': Comprador.Fuente.APOLLO,
    'procolombia': Comprador.Fuente.PROCOLOMBIA,
    'camara de comercio': Comprador.Fuente.CAMARA_COMERCIO,
    'cámara de comercio': Comprador.Fuente.CAMARA_COMERCIO,
    'linkedin': Comprador.Fuente.LINKEDIN,
    'linkedin sales navigator': Comprador.Fuente.LINKEDIN,
    'manual': Comprador.Fuente.MANUAL,
}


def _construir_mensaje(comprador, plantilla_email=_SIN_CARGAR, plantilla_wa=_SIN_CARGAR):
    """Devuelve (asunto_email, cuerpo_email, cuerpo_whatsapp) usando la primera
    plantilla disponible de cada tipo, o un mensaje genérico si no hay ninguna.

    Al recorrer muchos compradores (listados, envío masivo), pasa las
    plantillas ya buscadas una sola vez con `plantilla_email`/`plantilla_wa`
    — si no se pasan, esta función las busca ella misma (cómodo para un
    solo comprador, pero dispararía una consulta por fila en un bucle)."""
    producto = comprador.productos_interes.first()
    producto_nombre = producto.nombre if producto else ''

    if plantilla_email is _SIN_CARGAR:
        plantilla_email = PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.EMAIL).first()
    if plantilla_email:
        asunto, cuerpo_email = plantilla_email.render(comprador, producto_nombre)
    else:
        asunto = f'Oferta de producto para {comprador.nombre_empresa}'
        cuerpo_email = (
            f'Hola,\n\nNos gustaría ofrecerles nuestros productos disponibles en bodega'
            + (f' ({producto_nombre})' if producto_nombre else '')
            + '.\n\nQuedamos atentos.\n'
        )

    if plantilla_wa is _SIN_CARGAR:
        plantilla_wa = PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.WHATSAPP).first()
    if plantilla_wa:
        _, cuerpo_whatsapp = plantilla_wa.render(comprador, producto_nombre)
    else:
        cuerpo_whatsapp = (
            f'Hola {comprador.nombre_empresa}, les escribimos para ofrecerles nuestros productos'
            + (f' ({producto_nombre})' if producto_nombre else '')
            + '. ¿Les interesaría recibir más información?'
        )

    return asunto, cuerpo_email, cuerpo_whatsapp


def _aplicar_filtros_comprador(qs, params):
    """Aplica al queryset los mismos filtros de búsqueda (q/estado/pais/sector/fuente)
    usados tanto en el listado de compradores como en el envío masivo."""
    q = params.get('q', '').strip()
    estado = params.get('estado', '')
    pais = params.get('pais', '')
    sector = params.get('sector', '')
    fuente = params.get('fuente', '')

    if q:
        qs = qs.filter(
            Q(nombre_empresa__icontains=q) | Q(pais__icontains=q) | Q(sector__icontains=q)
        )
    if estado:
        qs = qs.filter(estado=estado)
    if pais:
        qs = qs.filter(pais=pais)
    if sector:
        qs = qs.filter(sector=sector)
    if fuente:
        qs = qs.filter(fuente=fuente)

    filtros = {'q': q, 'estado': estado, 'pais': pais, 'sector': sector, 'fuente': fuente}
    return qs, filtros


ETAPA_INDICE = {
    Comprador.Estado.POR_CONTACTAR: 0,
    Comprador.Estado.CONTACTADO: 1,
    Comprador.Estado.INTERESADO: 2,
    Comprador.Estado.CLIENTE: 3,
}


def _enriquecer(comprador, plantilla_email=_SIN_CARGAR, plantilla_wa=_SIN_CARGAR):
    asunto, cuerpo_email, cuerpo_whatsapp = _construir_mensaje(comprador, plantilla_email, plantilla_wa)
    comprador.mailto_url = f'mailto:{comprador.email}?subject={quote(asunto)}&body={quote(cuerpo_email)}'
    numero = comprador.whatsapp_numero
    comprador.whatsapp_url = f'https://wa.me/{numero}?text={quote(cuerpo_whatsapp)}' if numero else ''
    comprador.es_internacional = bool(comprador.pais) and comprador.pais != 'Colombia'
    comprador.etapa_indice = ETAPA_INDICE.get(comprador.estado)
    comprador.es_descartado = comprador.estado == Comprador.Estado.DESCARTADO
    return comprador


@login_required
def compradores_lista(request):
    qs = Comprador.objects.prefetch_related('productos_interes')
    qs, filtros = _aplicar_filtros_comprador(qs, request.GET)

    paginator = Paginator(qs, 20)
    page_obj = paginator.get_page(request.GET.get('page'))
    plantilla_email = PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.EMAIL).first()
    plantilla_wa = PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.WHATSAPP).first()
    for comprador in page_obj.object_list:
        _enriquecer(comprador, plantilla_email, plantilla_wa)

    paises = Comprador.objects.exclude(pais='').values_list('pais', flat=True).distinct().order_by('pais')
    sectores = Comprador.objects.exclude(sector='').values_list('sector', flat=True).distinct().order_by('sector')

    querystring = request.GET.copy()
    querystring.pop('page', None)

    querystring_sin_estado = querystring.copy()
    querystring_sin_estado.pop('estado', None)

    todos = Comprador.objects.all()
    total = todos.count()
    conteos = {
        valor: todos.filter(estado=valor).count()
        for valor, _ in Comprador.Estado.choices
    }
    embudo = []
    for valor, etiqueta in Comprador.Estado.choices:
        cantidad = conteos.get(valor, 0)
        embudo.append({
            'valor': valor,
            'etiqueta': etiqueta,
            'cantidad': cantidad,
            'pct': round(cantidad / total * 100) if total else 0,
            'activo': filtros['estado'] == valor,
        })

    context = {
        'page_obj': page_obj,
        'estados': Comprador.Estado.choices,
        'fuentes': Comprador.Fuente.choices,
        'paises': paises,
        'sectores': sectores,
        'filtros': filtros,
        'querystring': querystring.urlencode(),
        'querystring_sin_estado': querystring_sin_estado.urlencode(),
        'total_compradores': total,
        'embudo': embudo,
    }
    return render(request, 'prospeccion/compradores_lista.html', context)


@login_required
def comprador_detalle(request, pk):
    comprador = get_object_or_404(Comprador, pk=pk)
    _enriquecer(comprador)

    if request.method == 'POST':
        form = HistorialContactoForm(request.POST)
        if form.is_valid():
            historial = form.save(commit=False)
            historial.comprador = comprador
            historial.usuario = request.user
            historial.save()
            comprador.fecha_ultimo_contacto = timezone.now()
            if comprador.estado == Comprador.Estado.POR_CONTACTAR:
                comprador.estado = Comprador.Estado.CONTACTADO
            comprador.save()
            messages.success(request, 'Contacto registrado en el historial.')
            return redirect('comprador_detalle', pk=comprador.pk)
    else:
        form = HistorialContactoForm()

    context = {
        'comprador': comprador,
        'historial': comprador.historial.select_related('usuario'),
        'form': form,
    }
    return render(request, 'prospeccion/comprador_detalle.html', context)


@login_required
def comprador_form(request, pk=None):
    comprador = get_object_or_404(Comprador, pk=pk) if pk else None
    if request.method == 'POST':
        form = CompradorForm(request.POST, instance=comprador)
        if form.is_valid():
            form.save()
            messages.success(request, 'Comprador guardado correctamente.')
            return redirect('compradores_lista')
    else:
        form = CompradorForm(instance=comprador)
    return render(request, 'prospeccion/comprador_form.html', {'form': form, 'comprador': comprador})


@login_required
@require_POST
def comprador_eliminar(request, pk):
    comprador = get_object_or_404(Comprador, pk=pk)
    comprador.delete()
    messages.success(request, 'Comprador eliminado.')
    return redirect('compradores_lista')


@login_required
@require_POST
def comprador_marcar_contactado(request, pk):
    comprador = get_object_or_404(Comprador, pk=pk)
    comprador.fecha_ultimo_contacto = timezone.now()
    if comprador.estado == Comprador.Estado.POR_CONTACTAR:
        comprador.estado = Comprador.Estado.CONTACTADO
    comprador.save()
    HistorialContacto.objects.create(
        comprador=comprador,
        medio=request.POST.get('medio', HistorialContacto.Medio.OTRO),
        resultado='Marcado como contactado desde el listado.',
        usuario=request.user,
    )
    messages.success(request, f'{comprador.nombre_empresa} marcado como contactado.')
    next_url = request.POST.get('next') or reverse('compradores_lista')
    return redirect(next_url)


ENVIO_MASIVO_LIMITE = 300


@login_required
def compradores_envio_masivo(request):
    qs, filtros = _aplicar_filtros_comprador(Comprador.objects.all(), request.GET)
    qs = qs.exclude(email='').order_by('nombre_empresa').prefetch_related('productos_interes')
    total_filtrados = qs.count()
    destinatarios = list(qs[:ENVIO_MASIVO_LIMITE])

    plantilla_id = request.GET.get('plantilla', '')
    plantilla = None
    if plantilla_id:
        plantilla = PlantillaMensaje.objects.filter(pk=plantilla_id, tipo=PlantillaMensaje.Tipo.EMAIL).first()
    plantilla_email_generica = plantilla or PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.EMAIL).first()
    plantilla_wa = PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.WHATSAPP).first()
    for comprador in destinatarios:
        producto = comprador.productos_interes.first()
        producto_nombre = producto.nombre if producto else ''
        if plantilla:
            asunto, _ = plantilla.render(comprador, producto_nombre)
        else:
            asunto, _, _ = _construir_mensaje(
                comprador, plantilla_email=plantilla_email_generica, plantilla_wa=plantilla_wa,
            )
        comprador.asunto_preview = asunto
        comprador.whatsapp_url = f'https://wa.me/{comprador.whatsapp_numero}' if comprador.whatsapp_numero else ''

    context = {
        'destinatarios': destinatarios,
        'total_filtrados': total_filtrados,
        'limite': ENVIO_MASIVO_LIMITE,
        'estados': Comprador.Estado.choices,
        'fuentes': Comprador.Fuente.choices,
        'paises': Comprador.objects.exclude(pais='').values_list('pais', flat=True).distinct().order_by('pais'),
        'sectores': Comprador.objects.exclude(sector='').values_list('sector', flat=True).distinct().order_by('sector'),
        'filtros': filtros,
        'plantillas': PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.EMAIL),
        'plantilla_id': plantilla_id,
        'querystring': request.GET.urlencode(),
    }
    return render(request, 'prospeccion/compradores_envio_masivo.html', context)


@login_required
@require_POST
def compradores_envio_masivo_enviar(request):
    plantilla = get_object_or_404(
        PlantillaMensaje, pk=request.POST.get('plantilla_id'), tipo=PlantillaMensaje.Tipo.EMAIL,
    )
    pks = request.POST.getlist('compradores')
    compradores = Comprador.objects.filter(pk__in=pks).exclude(email='').prefetch_related('productos_interes')

    enviados = 0
    errores = 0
    connection = get_connection()
    connection.open()
    try:
        for comprador in compradores:
            producto = comprador.productos_interes.first()
            asunto, cuerpo = plantilla.render(comprador, producto.nombre if producto else '')
            try:
                EmailMessage(
                    subject=asunto, body=cuerpo, to=[comprador.email], connection=connection,
                ).send()
            except Exception:
                errores += 1
                continue

            HistorialContacto.objects.create(
                comprador=comprador,
                medio=HistorialContacto.Medio.EMAIL,
                resultado=f'Envío masivo — plantilla "{plantilla.nombre}"',
                usuario=request.user,
            )
            comprador.fecha_ultimo_contacto = timezone.now()
            if comprador.estado == Comprador.Estado.POR_CONTACTAR:
                comprador.estado = Comprador.Estado.CONTACTADO
            comprador.save()
            enviados += 1
    finally:
        connection.close()

    if enviados:
        messages.success(request, f'Se enviaron {enviados} correo(s) con la plantilla "{plantilla.nombre}".')
    if errores:
        messages.warning(request, f'{errores} correo(s) no se pudieron enviar.')
    if not enviados and not errores:
        messages.warning(request, 'No había destinatarios válidos seleccionados.')

    return redirect('compradores_lista')


@login_required
def compradores_importar(request):
    resumen = None
    if request.method == 'POST':
        form = ImportarCompradoresForm(request.POST, request.FILES)
        if form.is_valid():
            resumen = _procesar_importacion(form.cleaned_data['archivo'])
            if resumen['errores']:
                messages.warning(request, f"Importación completada con {len(resumen['errores'])} error(es).")
            else:
                messages.success(request, f"Se importaron {resumen['creados']} compradores nuevos.")
    else:
        form = ImportarCompradoresForm()
    return render(request, 'prospeccion/compradores_importar.html', {'form': form, 'resumen': resumen})


IMPORTACION_TAMANO_MAXIMO = 10 * 1024 * 1024  # 10 MB


def _procesar_importacion(archivo):
    nombre = archivo.name.lower()
    filas = []
    errores = []

    if archivo.size > IMPORTACION_TAMANO_MAXIMO:
        return {'creados': 0, 'errores': ['El archivo supera el tamaño máximo permitido (10 MB).']}

    if nombre.endswith('.csv'):
        contenido = archivo.read().decode('utf-8-sig', errors='replace')
        lector = csv.DictReader(io.StringIO(contenido))
        for fila in lector:
            filas.append({(k or '').strip().lower(): (v or '').strip() for k, v in fila.items()})
    elif nombre.endswith('.xlsx') or nombre.endswith('.xls'):
        import openpyxl

        wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
        ws = wb.active
        filas_iter = ws.iter_rows(values_only=True)
        encabezados = [str(h).strip().lower() if h else '' for h in next(filas_iter)]
        for fila in filas_iter:
            valores = {}
            for header, valor in zip(encabezados, fila):
                if header:
                    valores[header] = str(valor).strip() if valor is not None else ''
            if any(valores.values()):
                filas.append(valores)
    else:
        return {'creados': 0, 'errores': ['Formato de archivo no soportado. Usa CSV o Excel (.xlsx).']}

    creados = 0
    for i, fila in enumerate(filas, start=2):
        nombre_empresa = fila.get('nombre_empresa') or fila.get('empresa') or ''
        if not nombre_empresa:
            errores.append(f'Fila {i}: falta nombre_empresa, se omitió.')
            continue
        nombre_empresa = titulo_inteligente(nombre_empresa)

        fuente_raw = (fila.get('fuente') or '').strip().lower()
        fuente = FUENTE_ALIASES.get(fuente_raw, Comprador.Fuente.MANUAL)

        email = fila.get('email') or fila.get('correo') or ''
        existente = None
        if email:
            existente = Comprador.objects.filter(nombre_empresa__iexact=nombre_empresa, email__iexact=email).first()
        if existente:
            errores.append(f'Fila {i}: "{nombre_empresa}" ya existe, se omitió.')
            continue

        Comprador.objects.create(
            nombre_empresa=nombre_empresa,
            pais=fila.get('pais') or fila.get('país') or '',
            ciudad=fila.get('ciudad') or '',
            sector=fila.get('sector') or fila.get('industria') or '',
            email=email,
            telefono=fila.get('telefono') or fila.get('teléfono') or fila.get('whatsapp') or '',
            linkedin_url=fila.get('linkedin') or '',
            facebook_url=fila.get('facebook') or '',
            instagram_url=fila.get('instagram') or '',
            sitio_web=fila.get('sitio_web') or fila.get('web') or fila.get('website') or '',
            fuente=fuente,
            palabras_clave_interes=(
                fila.get('palabras_clave') or fila.get('palabras_clave_interes') or fila.get('intereses') or ''
            )[:300],
        )
        creados += 1

    return {'creados': creados, 'errores': errores}


@login_required
def productos_importar(request):
    resumen = None
    if request.method == 'POST':
        form = ImportarProductosForm(request.POST, request.FILES)
        if form.is_valid():
            resumen = _procesar_importacion_productos(form.cleaned_data['archivo'])
            if resumen['errores']:
                messages.warning(request, f"Importación completada con {len(resumen['errores'])} error(es).")
            else:
                messages.success(request, f"Se importaron {resumen['creados']} productos nuevos.")
    else:
        form = ImportarProductosForm()
    return render(request, 'prospeccion/productos_importar.html', {'form': form, 'resumen': resumen})


def _valor_booleano(texto):
    return (texto or '').strip().lower() in ('si', 'sí', 'true', '1', 'x', 'yes')


def _procesar_importacion_productos(archivo):
    nombre_archivo = archivo.name.lower()
    filas = []
    errores = []

    if archivo.size > IMPORTACION_TAMANO_MAXIMO:
        return {'creados': 0, 'errores': ['El archivo supera el tamaño máximo permitido (10 MB).']}

    if nombre_archivo.endswith('.csv'):
        contenido = archivo.read().decode('utf-8-sig', errors='replace')
        lector = csv.DictReader(io.StringIO(contenido))
        for fila in lector:
            filas.append({(k or '').strip().lower(): (v or '').strip() for k, v in fila.items()})
    elif nombre_archivo.endswith('.xlsx') or nombre_archivo.endswith('.xls'):
        import openpyxl

        wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
        ws = wb.active
        filas_iter = ws.iter_rows(values_only=True)
        encabezados = [str(h).strip().lower() if h else '' for h in next(filas_iter)]
        for fila in filas_iter:
            valores = {}
            for header, valor in zip(encabezados, fila):
                if header:
                    valores[header] = str(valor).strip() if valor is not None else ''
            if any(valores.values()):
                filas.append(valores)
    else:
        return {'creados': 0, 'errores': ['Formato de archivo no soportado. Usa CSV o Excel (.xlsx).']}

    creados = 0
    for i, fila in enumerate(filas, start=2):
        nombre = fila.get('nombre') or ''
        if not nombre:
            errores.append(f'Fila {i}: falta nombre, se omitió.')
            continue
        nombre = titulo_inteligente(nombre)

        referencia = fila.get('referencia') or ''
        if referencia and Producto.objects.filter(referencia__iexact=referencia).exists():
            errores.append(f'Fila {i}: la referencia "{referencia}" ya existe, se omitió (duplicado).')
            continue

        cantidad_raw = fila.get('cantidad_disponible') or ''
        try:
            cantidad_disponible = int(float(cantidad_raw)) if cantidad_raw else 0
            if cantidad_disponible < 0:
                raise ValueError
        except (TypeError, ValueError):
            if cantidad_raw:
                errores.append(f'Fila {i}: cantidad_disponible "{cantidad_raw}" no es válida, se puso en 0.')
            cantidad_disponible = 0

        def _decimal_o_advertencia(campo):
            crudo = fila.get(campo) or ''
            if not crudo:
                return None
            try:
                return Decimal(crudo.replace(',', ''))
            except InvalidOperation:
                errores.append(f'Fila {i}: {campo} "{crudo}" no es válido, se dejó vacío.')
                return None

        Producto.objects.create(
            nombre=nombre,
            categoria=fila.get('categoria') or '',
            descripcion=fila.get('descripcion') or '',
            cantidad_disponible=cantidad_disponible,
            precio_referencia=_decimal_o_advertencia('precio_referencia'),
            valor_estimado=_decimal_o_advertencia('valor_estimado'),
            condiciones_venta=fila.get('condiciones_venta') or '',
            referencia=referencia,
            confidencial=_valor_booleano(fila.get('confidencial')),
            proveedor_nombre=fila.get('proveedor_nombre') or '',
            proveedor_contacto=fila.get('proveedor_contacto') or '',
            proveedor_email=fila.get('proveedor_email') or '',
            proveedor_telefono=fila.get('proveedor_telefono') or '',
        )
        creados += 1

    return {'creados': creados, 'errores': errores}


# --- Búsqueda de compradores con IA ---
#
# La búsqueda con grounding real puede tardar decenas de segundos, así que
# no bloqueamos la petición HTTP esperando a la IA: se crea el registro de
# BusquedaIA de una vez (en estado "pendiente") y un hilo aparte hace la
# llamada real y actualiza ese mismo registro cuando termina. La pantalla
# de "Buscar con IA" solo guarda en sesión el ID de esa búsqueda y consulta
# su estado cada pocos segundos (ver compradores_buscar_ia_estado) hasta
# que queda "listo" o "error".

def _ejecutar_busqueda_en_segundo_plano(busqueda_id, consulta, funcion_busqueda=ia_busqueda.buscar_empresas):
    from django.db import connections

    try:
        empresas, fuentes = funcion_busqueda(consulta)
    except ia_busqueda.BusquedaIAError as exc:
        BusquedaIA.objects.filter(pk=busqueda_id).update(
            estado=BusquedaIA.Estado.ERROR, error_mensaje=str(exc),
        )
    except Exception:  # noqa: BLE001 — un hilo de fondo nunca debe morir en silencio sin dejar rastro
        BusquedaIA.objects.filter(pk=busqueda_id).update(
            estado=BusquedaIA.Estado.ERROR,
            error_mensaje='Ocurrió un error inesperado buscando con IA.',
        )
    else:
        BusquedaIA.objects.filter(pk=busqueda_id).update(
            estado=BusquedaIA.Estado.LISTO,
            resultados=len(empresas),
            resultados_json=empresas,
            fuentes_json=fuentes,
        )
    finally:
        connections.close_all()


@login_required
def compradores_buscar_ia(request):
    form = BusquedaIAForm(initial={'consulta': request.session.get('ia_consulta', '')})
    busqueda = None
    busqueda_id = request.session.get('ia_busqueda_id')
    if busqueda_id:
        busqueda = BusquedaIA.objects.filter(pk=busqueda_id, usuario=request.user).first()

    if busqueda and busqueda.estado == BusquedaIA.Estado.ERROR:
        messages.error(request, busqueda.error_mensaje or 'La IA no respondió correctamente.')
    elif busqueda and busqueda.estado == BusquedaIA.Estado.LISTO:
        if not busqueda.resultados_json:
            messages.warning(request, 'La IA no encontró empresas para esa búsqueda. Prueba con otros términos.')

    context = {
        'form': form,
        'busqueda': busqueda,
        'resultados': busqueda.resultados_json if busqueda and busqueda.estado == BusquedaIA.Estado.LISTO else None,
        'fuentes': (busqueda.fuentes_json or []) if busqueda and busqueda.estado == BusquedaIA.Estado.LISTO else [],
        'consulta_previa': request.session.get('ia_consulta', ''),
        'ia_configurada': bool(settings.GEMINI_API_KEY),
    }
    return render(request, 'prospeccion/compradores_buscar_ia.html', context)


@login_required
@require_POST
def compradores_buscar_ia_ejecutar(request):
    form = BusquedaIAForm(request.POST)
    if not form.is_valid():
        messages.error(request, 'Escribe qué quieres buscar.')
        return redirect('compradores_buscar_ia')

    consulta = form.cleaned_data['consulta']

    hoy = timezone.localdate()
    busquedas_hoy = BusquedaIA.objects.filter(creado_en__date=hoy).count()
    if busquedas_hoy >= settings.BUSQUEDA_IA_LIMITE_DIARIO:
        messages.error(
            request,
            f'Ya se hicieron {busquedas_hoy} búsquedas con IA hoy '
            f'(tope diario: {settings.BUSQUEDA_IA_LIMITE_DIARIO}). Intenta de nuevo mañana.',
        )
        return redirect('compradores_buscar_ia')

    busqueda = BusquedaIA.objects.create(
        consulta=consulta, usuario=request.user, estado=BusquedaIA.Estado.PENDIENTE,
    )
    threading.Thread(
        target=_ejecutar_busqueda_en_segundo_plano, args=(busqueda.pk, consulta), daemon=True,
    ).start()

    request.session['ia_busqueda_id'] = busqueda.pk
    request.session['ia_consulta'] = consulta
    return redirect('compradores_buscar_ia')


@login_required
def compradores_buscar_ia_estado(request, busqueda_id):
    busqueda = get_object_or_404(BusquedaIA, pk=busqueda_id, usuario=request.user)
    return JsonResponse({
        'estado': busqueda.estado,
        'total_resultados': busqueda.resultados if busqueda.estado == BusquedaIA.Estado.LISTO else None,
    })


@login_required
@require_POST
def compradores_buscar_ia_guardar(request):
    busqueda_id = request.session.get('ia_busqueda_id')
    busqueda = BusquedaIA.objects.filter(pk=busqueda_id, usuario=request.user).first() if busqueda_id else None
    resultados = (busqueda.resultados_json or []) if busqueda else []
    seleccionados = request.POST.getlist('seleccion')

    creados = 0
    for indice in seleccionados:
        try:
            item = resultados[int(indice)]
        except (ValueError, IndexError, TypeError):
            continue

        nombre_empresa = item.get('nombre_empresa')
        if not nombre_empresa:
            continue

        existente = None
        if item.get('email'):
            existente = Comprador.objects.filter(
                nombre_empresa__iexact=nombre_empresa, email__iexact=item['email'],
            ).first()
        if existente:
            continue

        notas = 'Encontrado con búsqueda de IA — verificar antes de contactar.'
        if item.get('resumen'):
            notas += f"\n{item['resumen']}"

        Comprador.objects.create(
            nombre_empresa=nombre_empresa,
            pais=item.get('pais', ''),
            ciudad=item.get('ciudad', ''),
            sector=item.get('sector', ''),
            email=item.get('email', ''),
            telefono=item.get('telefono', ''),
            sitio_web=item.get('sitio_web', ''),
            linkedin_url=item.get('linkedin_url', ''),
            facebook_url=item.get('facebook_url', ''),
            instagram_url=item.get('instagram_url', ''),
            fuente=Comprador.Fuente.IA,
            notas=notas,
        )
        creados += 1

    request.session.pop('ia_busqueda_id', None)
    request.session.pop('ia_consulta', None)

    messages.success(request, f'Se guardaron {creados} compradores nuevos.')
    return redirect('compradores_lista')


@login_required
@require_POST
def compradores_buscar_ia_descartar(request):
    request.session.pop('ia_busqueda_id', None)
    request.session.pop('ia_consulta', None)
    return redirect('compradores_buscar_ia')


# --- Proveedores / consignantes ---
# El lado opuesto de Compradores: empresas que le dan inventario a ISYN para
# vender/subastar (incluye chatarreos). Reutiliza el mismo patrón de listado,
# ficha, formulario e historial, y la misma infraestructura de búsqueda con
# IA en segundo plano — solo cambia el modelo y el prompt.

ETAPA_INDICE_PROVEEDOR = {
    Proveedor.Estado.POR_CONTACTAR: 0,
    Proveedor.Estado.CONTACTADO: 1,
    Proveedor.Estado.NEGOCIANDO: 2,
    Proveedor.Estado.CONSIGNO: 3,
}


def _aplicar_filtros_proveedor(qs, params):
    q = params.get('q', '').strip()
    estado = params.get('estado', '')
    pais = params.get('pais', '')
    sector = params.get('sector', '')
    fuente = params.get('fuente', '')

    if q:
        qs = qs.filter(
            Q(nombre_empresa__icontains=q) | Q(pais__icontains=q) | Q(sector__icontains=q)
        )
    if estado:
        qs = qs.filter(estado=estado)
    if pais:
        qs = qs.filter(pais=pais)
    if sector:
        qs = qs.filter(sector=sector)
    if fuente:
        qs = qs.filter(fuente=fuente)

    filtros = {'q': q, 'estado': estado, 'pais': pais, 'sector': sector, 'fuente': fuente}
    return qs, filtros


def _enriquecer_proveedor(proveedor):
    asunto = 'Servicios de subasta y venta — ISYN'
    cuerpo = (
        f'Hola {proveedor.nombre_empresa},\n\nEn ISYN ayudamos a empresas como la suya a vender o '
        'subastar activos e inventario. Nos gustaría conversar sobre cómo podemos ayudarles.\n\nQuedamos atentos.\n'
    )
    proveedor.mailto_url = f'mailto:{proveedor.email}?subject={quote(asunto)}&body={quote(cuerpo)}'
    numero = proveedor.whatsapp_numero
    proveedor.whatsapp_url = f'https://wa.me/{numero}?text={quote(cuerpo)}' if numero else ''
    proveedor.es_internacional = bool(proveedor.pais) and proveedor.pais != 'Colombia'
    proveedor.etapa_indice = ETAPA_INDICE_PROVEEDOR.get(proveedor.estado)
    proveedor.es_descartado = proveedor.estado == Proveedor.Estado.DESCARTADO
    return proveedor


@login_required
def proveedores_lista(request):
    qs, filtros = _aplicar_filtros_proveedor(Proveedor.objects.all(), request.GET)

    paginator = Paginator(qs, 20)
    page_obj = paginator.get_page(request.GET.get('page'))
    for proveedor in page_obj.object_list:
        _enriquecer_proveedor(proveedor)

    paises = Proveedor.objects.exclude(pais='').values_list('pais', flat=True).distinct().order_by('pais')
    sectores = Proveedor.objects.exclude(sector='').values_list('sector', flat=True).distinct().order_by('sector')

    querystring = request.GET.copy()
    querystring.pop('page', None)

    querystring_sin_estado = querystring.copy()
    querystring_sin_estado.pop('estado', None)

    todos = Proveedor.objects.all()
    total = todos.count()
    conteos = {
        valor: todos.filter(estado=valor).count()
        for valor, _ in Proveedor.Estado.choices
    }
    embudo = []
    for valor, etiqueta in Proveedor.Estado.choices:
        cantidad = conteos.get(valor, 0)
        embudo.append({
            'valor': valor,
            'etiqueta': etiqueta,
            'cantidad': cantidad,
            'pct': round(cantidad / total * 100) if total else 0,
            'activo': filtros['estado'] == valor,
        })

    context = {
        'page_obj': page_obj,
        'estados': Proveedor.Estado.choices,
        'fuentes': Proveedor.Fuente.choices,
        'paises': paises,
        'sectores': sectores,
        'filtros': filtros,
        'querystring': querystring.urlencode(),
        'querystring_sin_estado': querystring_sin_estado.urlencode(),
        'total_proveedores': total,
        'embudo': embudo,
    }
    return render(request, 'prospeccion/proveedores_lista.html', context)


@login_required
def proveedor_detalle(request, pk):
    proveedor = get_object_or_404(Proveedor, pk=pk)
    _enriquecer_proveedor(proveedor)

    if request.method == 'POST':
        form = HistorialContactoProveedorForm(request.POST)
        if form.is_valid():
            historial = form.save(commit=False)
            historial.proveedor = proveedor
            historial.usuario = request.user
            historial.save()
            proveedor.fecha_ultimo_contacto = timezone.now()
            if proveedor.estado == Proveedor.Estado.POR_CONTACTAR:
                proveedor.estado = Proveedor.Estado.CONTACTADO
            proveedor.save()
            messages.success(request, 'Contacto registrado en el historial.')
            return redirect('proveedor_detalle', pk=proveedor.pk)
    else:
        form = HistorialContactoProveedorForm()

    context = {
        'proveedor': proveedor,
        'historial': proveedor.historial.select_related('usuario'),
        'form': form,
    }
    return render(request, 'prospeccion/proveedor_detalle.html', context)


@login_required
def proveedor_form(request, pk=None):
    proveedor = get_object_or_404(Proveedor, pk=pk) if pk else None
    if request.method == 'POST':
        form = ProveedorForm(request.POST, instance=proveedor)
        if form.is_valid():
            form.save()
            messages.success(request, 'Proveedor guardado correctamente.')
            return redirect('proveedores_lista')
    else:
        form = ProveedorForm(instance=proveedor)
    return render(request, 'prospeccion/proveedor_form.html', {'form': form, 'proveedor': proveedor})


@login_required
@require_POST
def proveedor_eliminar(request, pk):
    proveedor = get_object_or_404(Proveedor, pk=pk)
    proveedor.delete()
    messages.success(request, 'Proveedor eliminado.')
    return redirect('proveedores_lista')


@login_required
@require_POST
def proveedor_marcar_contactado(request, pk):
    proveedor = get_object_or_404(Proveedor, pk=pk)
    proveedor.fecha_ultimo_contacto = timezone.now()
    if proveedor.estado == Proveedor.Estado.POR_CONTACTAR:
        proveedor.estado = Proveedor.Estado.CONTACTADO
    proveedor.save()
    HistorialContactoProveedor.objects.create(
        proveedor=proveedor,
        medio=request.POST.get('medio', HistorialContacto.Medio.OTRO),
        resultado='Marcado como contactado desde el listado.',
        usuario=request.user,
    )
    messages.success(request, f'{proveedor.nombre_empresa} marcado como contactado.')
    next_url = request.POST.get('next') or reverse('proveedores_lista')
    return redirect(next_url)


@login_required
def proveedores_buscar_ia(request):
    form = BusquedaIAForm(initial={'consulta': request.session.get('ia_consulta_proveedor', '')})
    busqueda = None
    busqueda_id = request.session.get('ia_busqueda_proveedor_id')
    if busqueda_id:
        busqueda = BusquedaIA.objects.filter(
            pk=busqueda_id, usuario=request.user, tipo=BusquedaIA.Tipo.PROVEEDOR,
        ).first()

    if busqueda and busqueda.estado == BusquedaIA.Estado.ERROR:
        messages.error(request, busqueda.error_mensaje or 'La IA no respondió correctamente.')
    elif busqueda and busqueda.estado == BusquedaIA.Estado.LISTO:
        if not busqueda.resultados_json:
            messages.warning(request, 'La IA no encontró empresas para esa búsqueda. Prueba con otros términos.')

    context = {
        'form': form,
        'busqueda': busqueda,
        'resultados': busqueda.resultados_json if busqueda and busqueda.estado == BusquedaIA.Estado.LISTO else None,
        'fuentes': (busqueda.fuentes_json or []) if busqueda and busqueda.estado == BusquedaIA.Estado.LISTO else [],
        'consulta_previa': request.session.get('ia_consulta_proveedor', ''),
        'ia_configurada': bool(settings.GEMINI_API_KEY),
    }
    return render(request, 'prospeccion/proveedores_buscar_ia.html', context)


@login_required
@require_POST
def proveedores_buscar_ia_ejecutar(request):
    form = BusquedaIAForm(request.POST)
    if not form.is_valid():
        messages.error(request, 'Escribe qué quieres buscar.')
        return redirect('proveedores_buscar_ia')

    consulta = form.cleaned_data['consulta']

    hoy = timezone.localdate()
    busquedas_hoy = BusquedaIA.objects.filter(creado_en__date=hoy).count()
    if busquedas_hoy >= settings.BUSQUEDA_IA_LIMITE_DIARIO:
        messages.error(
            request,
            f'Ya se hicieron {busquedas_hoy} búsquedas con IA hoy '
            f'(tope diario: {settings.BUSQUEDA_IA_LIMITE_DIARIO}). Intenta de nuevo mañana.',
        )
        return redirect('proveedores_buscar_ia')

    busqueda = BusquedaIA.objects.create(
        consulta=consulta, usuario=request.user, estado=BusquedaIA.Estado.PENDIENTE,
        tipo=BusquedaIA.Tipo.PROVEEDOR,
    )
    threading.Thread(
        target=_ejecutar_busqueda_en_segundo_plano,
        args=(busqueda.pk, consulta, ia_busqueda.buscar_proveedores),
        daemon=True,
    ).start()

    request.session['ia_busqueda_proveedor_id'] = busqueda.pk
    request.session['ia_consulta_proveedor'] = consulta
    return redirect('proveedores_buscar_ia')


@login_required
def proveedores_buscar_ia_estado(request, busqueda_id):
    busqueda = get_object_or_404(BusquedaIA, pk=busqueda_id, usuario=request.user, tipo=BusquedaIA.Tipo.PROVEEDOR)
    return JsonResponse({
        'estado': busqueda.estado,
        'total_resultados': busqueda.resultados if busqueda.estado == BusquedaIA.Estado.LISTO else None,
    })


@login_required
@require_POST
def proveedores_buscar_ia_guardar(request):
    busqueda_id = request.session.get('ia_busqueda_proveedor_id')
    busqueda = (
        BusquedaIA.objects.filter(pk=busqueda_id, usuario=request.user, tipo=BusquedaIA.Tipo.PROVEEDOR).first()
        if busqueda_id else None
    )
    resultados = (busqueda.resultados_json or []) if busqueda else []
    seleccionados = request.POST.getlist('seleccion')

    creados = 0
    for indice in seleccionados:
        try:
            item = resultados[int(indice)]
        except (ValueError, IndexError, TypeError):
            continue

        nombre_empresa = item.get('nombre_empresa')
        if not nombre_empresa:
            continue

        existente = None
        if item.get('email'):
            existente = Proveedor.objects.filter(
                nombre_empresa__iexact=nombre_empresa, email__iexact=item['email'],
            ).first()
        if existente:
            continue

        notas = 'Encontrado con búsqueda de IA — verificar antes de contactar.'
        if item.get('resumen'):
            notas += f"\n{item['resumen']}"

        Proveedor.objects.create(
            nombre_empresa=nombre_empresa,
            pais=item.get('pais', ''),
            ciudad=item.get('ciudad', ''),
            sector=item.get('sector', ''),
            email=item.get('email', ''),
            telefono=item.get('telefono', ''),
            sitio_web=item.get('sitio_web', ''),
            linkedin_url=item.get('linkedin_url', ''),
            facebook_url=item.get('facebook_url', ''),
            instagram_url=item.get('instagram_url', ''),
            fuente=Proveedor.Fuente.IA,
            notas=notas,
        )
        creados += 1

    request.session.pop('ia_busqueda_proveedor_id', None)
    request.session.pop('ia_consulta_proveedor', None)

    messages.success(request, f'Se guardaron {creados} proveedores nuevos.')
    return redirect('proveedores_lista')


@login_required
@require_POST
def proveedores_buscar_ia_descartar(request):
    request.session.pop('ia_busqueda_proveedor_id', None)
    request.session.pop('ia_consulta_proveedor', None)
    return redirect('proveedores_buscar_ia')


# --- Productos ---

@login_required
def productos_lista(request):
    productos = Producto.objects.all()
    q = request.GET.get('q', '').strip()
    categoria = request.GET.get('categoria', '')
    if q:
        productos = productos.filter(
            Q(nombre__icontains=q) | Q(categoria__icontains=q) | Q(referencia__icontains=q)
        )
    if categoria:
        productos = productos.filter(categoria=categoria)

    categorias = Producto.objects.exclude(categoria='').values_list('categoria', flat=True).distinct().order_by('categoria')

    paginator = Paginator(productos, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    querystring = request.GET.copy()
    querystring.pop('page', None)

    context = {
        'page_obj': page_obj,
        'total_productos': paginator.count,
        'q': q,
        'categoria': categoria,
        'categorias': categorias,
        'querystring': querystring.urlencode(),
    }
    return render(request, 'prospeccion/productos_lista.html', context)


@login_required
def productos_dashboard(request):
    productos = Producto.objects.all()
    total_productos = productos.count()
    cantidad_total = productos.aggregate(total=Sum('cantidad_disponible'))['total'] or 0
    valor_total = (productos.aggregate(v=Sum('valor_estimado'))['v'] or 0) + (
        productos.filter(valor_estimado__isnull=True)
        .aggregate(v=Sum('precio_referencia'))['v'] or 0
    )
    categorias_count = productos.exclude(categoria='').values('categoria').distinct().count()
    confidenciales_count = productos.filter(confidencial=True).count()

    categorias = list(
        productos.exclude(categoria='')
        .values('categoria')
        .annotate(total=Count('id'), cantidad=Sum('cantidad_disponible'))
        .order_by('-total')[:10]
    )
    max_total = max((c['total'] for c in categorias), default=0)
    for c in categorias:
        c['pct'] = round(c['total'] / max_total * 100) if max_total else 0

    top_interes = (
        Producto.objects.annotate(num_interesados=Count('compradores_interesados'))
        .filter(num_interesados__gt=0)
        .order_by('-num_interesados')[:8]
    )

    context = {
        'total_productos': total_productos,
        'cantidad_total': cantidad_total,
        'valor_total': valor_total,
        'categorias_count': categorias_count,
        'confidenciales_count': confidenciales_count,
        'categorias': categorias,
        'top_interes': top_interes,
    }
    return render(request, 'prospeccion/productos_dashboard.html', context)


@login_required
def producto_detalle(request, pk):
    producto = get_object_or_404(Producto, pk=pk)

    if request.method == 'POST':
        comprador = Comprador.objects.filter(pk=request.POST.get('comprador_id')).first()
        if comprador:
            producto.compradores_interesados.add(comprador)
            messages.success(request, f'{comprador.nombre_empresa} se vinculó como interesado en este producto.')
        return redirect('producto_detalle', pk=producto.pk)

    plantilla_email = PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.EMAIL).first()
    plantilla_wa = PlantillaMensaje.objects.filter(tipo=PlantillaMensaje.Tipo.WHATSAPP).first()
    interesados = [
        _enriquecer(c, plantilla_email, plantilla_wa)
        for c in producto.compradores_interesados.prefetch_related('productos_interes')
    ]
    candidatos = Comprador.objects.exclude(
        pk__in=[c.pk for c in interesados]
    ).order_by('nombre_empresa')

    context = {
        'producto': producto,
        'interesados': interesados,
        'candidatos': candidatos,
    }
    return render(request, 'prospeccion/producto_detalle.html', context)


@login_required
def producto_form(request, pk=None):
    producto = get_object_or_404(Producto, pk=pk) if pk else None
    if request.method == 'POST':
        form = ProductoForm(request.POST, instance=producto)
        if form.is_valid():
            form.save()
            messages.success(request, 'Producto guardado correctamente.')
            return redirect('productos_lista')
    else:
        form = ProductoForm(instance=producto)
    return render(request, 'prospeccion/producto_form.html', {'form': form, 'producto': producto})


@login_required
@require_POST
def producto_eliminar(request, pk):
    producto = get_object_or_404(Producto, pk=pk)
    producto.delete()
    messages.success(request, 'Producto eliminado.')
    return redirect('productos_lista')


# --- Plantillas ---

@login_required
def plantillas_lista(request):
    plantillas = PlantillaMensaje.objects.all()
    return render(request, 'prospeccion/plantillas_lista.html', {'plantillas': plantillas})


@login_required
def plantilla_form(request, pk=None):
    plantilla = get_object_or_404(PlantillaMensaje, pk=pk) if pk else None
    if request.method == 'POST':
        form = PlantillaMensajeForm(request.POST, instance=plantilla)
        if form.is_valid():
            form.save()
            messages.success(request, 'Plantilla guardada correctamente.')
            return redirect('plantillas_lista')
    else:
        form = PlantillaMensajeForm(instance=plantilla)
    return render(request, 'prospeccion/plantilla_form.html', {'form': form, 'plantilla': plantilla})


@login_required
@require_POST
def plantilla_eliminar(request, pk):
    plantilla = get_object_or_404(PlantillaMensaje, pk=pk)
    plantilla.delete()
    messages.success(request, 'Plantilla eliminada.')
    return redirect('plantillas_lista')


# --- Integración externa (preparado para futuro, ej. Hunter.io) ---

@login_required
@require_POST
def api_importar_hunter(request):
    """
    Endpoint preparado para integrar una API externa de prospección
    (ej. Hunter.io Domain Search) y guardar los resultados como nuevos
    Compradores con fuente=hunter.

    Aún no implementado: requiere configurar HUNTER_API_KEY en el entorno
    y añadir la llamada real a la API dentro de esta vista. Por ahora
    devuelve 501 para dejar el contrato de la ruta ya definido.

    Uso esperado (cuando esté implementado):
      POST /api/importar/hunter/  {"dominio": "empresa.com", "sector": "textil"}
      -> crea Compradores con fuente=Comprador.Fuente.HUNTER y devuelve
         {"creados": N, "compradores": [...]}
    """
    return JsonResponse(
        {'detail': 'Integración con Hunter.io aún no configurada.'},
        status=501,
    )
